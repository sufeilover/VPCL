"""Broad AC AI driving-characteristic preprocessing (Python 3.10+).

Reads only <root>/*+AI/<AIpush>/*.xlsx, one workbook at a time. Raw inputs are
never modified. No envelope, accuracy threshold, incident or causal claims.
Use --inspect to read only headers and a few rows; --self-test reads no inputs.
Dependencies: numpy, pandas, openpyxl.
"""
from __future__ import annotations
import argparse
import hashlib
import itertools
import json
import math
import os
import sys
import traceback
from collections import defaultdict
from pathlib import Path
import numpy as np
import pandas as pd

VERSION = '1.0.0'
ROOT_DEFAULT = Path('E:/IEEE-TVT/new-similar')
WHEELS = ('FL', 'FR', 'RL', 'RR')
ATTITUDE = ('heading', 'pitch', 'roll')
META = {'t_sec', 'packetId', 'udp_pid_used', 'completedLaps', 'normalizedCarPosition',
        'currentSectorIndex', 'lastSectorTime', 'iCurrentTime_s', 'iLastTime_s', 'iBestTime_s'}
DISCRETE = {'packetId', 'udp_pid_used', 'completedLaps', 'currentSectorIndex', 'gear', 'Gear'}
ANGLE = {'heading': 2*np.pi, 'pitch': 2*np.pi, 'roll': 2*np.pi}
PERCENTILES = (.01, .05, .10, .25, .50, .75, .90, .95, .99)

def arguments(argv=None):
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument('--root', type=Path, default=ROOT_DEFAULT)
    p.add_argument('--output', type=Path, help='Default: <root>/ai_driving_characteristics')
    p.add_argument('--settings', nargs='+', type=int, default=[80,85,90,95,100])
    p.add_argument('--sheet', default='telemetry')
    p.add_argument('--limit', type=int, help='Process only the first N discovered files (partial results labeled)')
    p.add_argument('--inspect', action='store_true', help='At most one file per configuration, first 3 rows only')
    p.add_argument('--self-test', action='store_true')
    p.add_argument('--expected-repeats', type=int, default=5)
    p.add_argument('--expected-configs', type=int, default=4)
    p.add_argument('--grid-hz', type=float, default=50.)
    p.add_argument('--smooth-seconds', type=float, default=.10)
    p.add_argument('--max-gap-seconds', type=float, default=.10, help='No interpolation/derivative across a larger time gap')
    p.add_argument('--profile-points', type=int, default=1000)
    p.add_argument('--max-progress-gap', type=float, default=.01, help='Do not interpolate spatial profiles across larger progress gaps')
    p.add_argument('--position-jump-m', type=float, default=2., help='QC flag only; not a crash threshold')
    p.add_argument('--max-progress-step', type=float, default=.10, help='Reject ambiguous lap extraction if a single progress increment exceeds this magnitude')
    p.add_argument('--direction', choices=['forward','reverse'], default='forward')
    p.add_argument('--attitude-unit', choices=['rad','deg'], default='rad')
    p.add_argument('--slip-angle-unit', choices=['deg','rad','native'], default='deg')
    p.add_argument('--steer-unit', choices=['native','rad','deg'], default='native')
    p.add_argument('--save-clean-laps', action='store_true', help='Optional all-column complete laps, compressed CSV')
    p.add_argument('--save-time-series', action='store_true', help='Optional smoothed/resampled and derived time series')
    p.add_argument('--fail-fast', action='store_true')
    p.add_argument('--no-hash', action='store_true', help='Faster, but provenance then uses only path/size/mtime')
    return p.parse_args(argv)

def sha(path):
    h=hashlib.sha256()
    with Path(path).open('rb') as f:
        for chunk in iter(lambda:f.read(1024*1024),b''): h.update(chunk)
    return h.hexdigest()

def write_csv(data, path):
    path=Path(path); path.parent.mkdir(parents=True,exist_ok=True)
    tmp=path.with_name(path.name+'.tmp')
    pd.DataFrame(data).to_csv(tmp,index=False,encoding='utf-8-sig',compression='gzip' if path.suffix=='.gz' else None)
    os.replace(tmp,path)

def write_json(data,path):
    path=Path(path);path.parent.mkdir(parents=True,exist_ok=True)
    tmp=path.with_name(path.name+'.tmp')
    tmp.write_text(json.dumps(data,ensure_ascii=False,indent=2,default=str,allow_nan=False),encoding='utf-8')
    os.replace(tmp,path)

def discover(a):
    found=[]
    for folder in sorted(a.root.glob('*+AI')):
        if not folder.is_dir(): continue
        config=folder.name[:-3]
        if '+' not in config: continue
        track,vehicle=config.split('+',1)
        for setting in sorted(set(a.settings)):
            for path in sorted((folder/str(setting)).glob('*.xlsx')):
                if not path.name.startswith('~$'):
                    found.append(dict(config=config,track=track,vehicle=vehicle,setting=setting,run=path.stem,source=str(path.resolve())))
    return found

def inspect(a,inputs):
    from openpyxl import load_workbook
    seen=set()
    for info in inputs:
        if info['config'] in seen: continue
        seen.add(info['config']); w=load_workbook(info['source'],read_only=True,data_only=True)
        try:
            if a.sheet not in w.sheetnames: raise ValueError(f'Missing sheet {a.sheet}')
            print(json.dumps({'source':info['source'],'sheets':w.sheetnames,'first_rows':list(w[a.sheet].iter_rows(max_row=4,values_only=True))},ensure_ascii=False,default=str))
        finally:w.close()

def read_and_extract(path,a):
    d=pd.read_excel(path,sheet_name=a.sheet,engine='openpyxl')
    d.columns=[str(c).strip() for c in d.columns]
    if d.columns.duplicated().any(): raise ValueError('Duplicate column names after trimming')
    if not {'t_sec','normalizedCarPosition'}.issubset(d): raise ValueError('Required columns: t_sec, normalizedCarPosition')
    qc={'original_rows':len(d)}
    inventory=[]
    for c in d:
        numeric=pd.to_numeric(d[c],errors='coerce').replace([np.inf,-np.inf],np.nan)
        inventory.append({'channel':c,'rows':len(d),'numeric_finite_rows':int(numeric.notna().sum()),
            'nonempty_rows':int(d[c].notna().sum()),'excluded_as_metadata':c in META})
        d[c]=numeric
    core=np.isfinite(d[['t_sec','normalizedCarPosition']]).all(axis=1)
    qc['dropped_invalid_time_or_progress']=int((~core).sum());d=d.loc[core].copy()
    if len(d)<4:raise ValueError('Fewer than four usable time/progress rows')
    if np.any(np.diff(d.t_sec)<0):raise ValueError('Timestamp regression/reset: file not sorted or silently repaired')
    qc['duplicate_timestamp_rows']=int(d.t_sec.duplicated().sum())
    d=d.drop_duplicates('t_sec',keep='last').reset_index(drop=True)
    p=d.normalizedCarPosition.to_numpy(float)
    if np.any((p<0)|(p>1)):raise ValueError('normalizedCarPosition outside [0,1]')
    dp=(np.diff(p)+.5)%1-.5
    if np.any(np.abs(dp)>a.max_progress_step):raise ValueError('Ambiguous large progress increment; inspect reset/gap before lap extraction')
    direction=1 if a.direction=='forward' else -1
    phase=np.r_[0.,np.cumsum(direction*dp)]
    hits=np.flatnonzero(phase>=1.)
    if len(hits)==0:raise ValueError(f'No complete circuit: maximum forward progress={phase.max():.6f}')
    end=int(hits[0]); lap=d.iloc[:end+1].copy()
    frac=(1-phase[end-1])/(phase[end]-phase[end-1])
    t=d.t_sec.to_numpy(float); exact_end=t[end-1]+frac*(t[end]-t[end-1])
    # Exact progress endpoint. Missing endpoints remain missing; discrete fields use left sample.
    for c in lap:
        left=float(d[c].iloc[end-1]);right=float(d[c].iloc[end])
        period=(360. if a.attitude_unit=='deg' else 2*np.pi) if c in ATTITUDE else None
        if c in DISCRETE: val=left
        elif np.isfinite(left) and np.isfinite(right):
            delta=((right-left+period/2)%period-period/2) if period else right-left
            val=left+frac*delta
        else:val=np.nan
        lap.loc[end,c]=val
    lap.loc[end,'t_sec']=exact_end
    lap.loc[end,'normalizedCarPosition']=(p[0]+direction)%1
    lap['_progress']=np.r_[phase[:end],1.]
    dt=np.diff(lap.t_sec)
    qc.update({'usable_rows':len(d),'lap_rows':len(lap),'direction':direction,'start_progress':float(p[0]),
        'lap_time_s':float(exact_end-t[0]),'sampled_crossing_lap_time_s':float(t[end]-t[0]),
        'endpoint_interpolation_correction_s':float(t[end]-exact_end),'post_lap_recording_s':float(t[-1]-exact_end),
        'median_dt_s':float(np.median(dt)),'p95_dt_s':float(np.quantile(dt,.95)),'max_dt_s':float(dt.max()),
        'time_gap_count':int((dt>a.max_gap_seconds).sum()),'time_gap_duration_s':float(dt[dt>a.max_gap_seconds].sum()),
        'progress_regression_rows':int((np.diff(lap['_progress'])< -1e-10).sum()),
        'max_progress_drawdown':float(np.max(np.maximum.accumulate(lap['_progress'])-lap['_progress']))})
    if {'x','y','z'}.issubset(lap):
        jumps=np.linalg.norm(np.diff(lap[['x','y','z']].to_numpy(float),axis=0),axis=1)
        finite=jumps[np.isfinite(jumps)]
        qc['position_jump_count']=int((finite>a.position_jump_m).sum())
        qc['max_position_jump_m']=float(finite.max()) if len(finite) else None
    for c in ['gas','brake']:
        if c in lap:
            y=lap[c].to_numpy(float)
            qc[c+'_outside_0_1_rows']=int(np.sum(np.isfinite(y)&((y< -1e-5)|(y>1+1e-5))))
    if {'packetId','udp_pid_used'}.issubset(lap):
        delta=lap.packetId-lap.udp_pid_used
        qc['pid_difference_mean']=float(delta.mean()) if delta.notna().any() else None
        qc['pid_difference_max']=float(delta.max()) if delta.notna().any() else None
    return lap,qc,inventory

def intervals(t,y,max_gap):
    dt=np.diff(t);good=np.isfinite(y[:-1]) & np.isfinite(y[1:]) & (dt>0) & (dt<=max_gap)
    return dt,good

def weighted_quantiles(y,w):
    good=np.isfinite(y)&(w>0);yy=y[good];ww=w[good]
    if not len(yy):return [np.nan]*len(PERCENTILES)
    order=np.argsort(yy);yy=yy[order];ww=ww[order]
    cdf=(np.cumsum(ww)-.5*ww)/ww.sum()
    return np.interp(PERCENTILES,cdf,yy)

def summarize(t,y,max_gap,period=None):
    y=np.asarray(y,float);dt,ok=intervals(t,y,max_gap)
    w=np.zeros(len(t));w[:-1]+=np.where(ok,dt/2,0);w[1:]+=np.where(ok,dt/2,0)
    duration=float(np.sum(dt[ok])); finite=y[np.isfinite(y)]
    out={'finite_fraction':float(np.isfinite(y).mean()),'valid_duration_s':duration,'coverage_fraction':duration/(t[-1]-t[0]),'unique_values':float(len(np.unique(finite)))}
    if not len(finite) or duration==0:return out
    if period:
        angle=y*2*np.pi/period;z=np.exp(1j*np.where(w>0,angle,0))
        z=np.sum(w*z)/w.sum();center=float(np.angle(z)*period/(2*np.pi))
        residual=(y-center+period/2)%period-period/2
        out.update({'circular_mean':center,'circular_resultant':float(abs(z)),
                    'circular_sd':float(np.sqrt(-2*np.log(np.clip(abs(z),1e-15,1)))*period/(2*np.pi))})
        out.update({f'centered_q{int(q*100):02d}':float(v) for q,v in zip(PERCENTILES,weighted_quantiles(residual,w))})
        return out
    safe=np.where(w>0,y,0);mean=float(np.sum(w*safe)/w.sum())
    variance=float(np.sum(w*(safe-mean)**2)/w.sum())
    out.update({'time_mean':mean,'time_sd':math.sqrt(max(variance,0)), 'time_rms':float(np.sqrt(np.sum(w*safe**2)/w.sum())),
        'time_mean_abs':float(np.sum(w*np.abs(safe))/w.sum()),'min':float(finite.min()),'max':float(finite.max()),
        'abs_max':float(np.max(np.abs(finite))),'range':float(np.ptp(finite)),
        'time_integral':float(np.sum(w*safe)), 'abs_time_integral':float(np.sum(w*np.abs(safe))),
        'total_variation':float(np.sum(np.abs(np.diff(y))[ok])),
        'variation_per_s':float(np.sum(np.abs(np.diff(y))[ok])/duration)})
    out.update({f'q{int(q*100):02d}':float(v) for q,v in zip(PERCENTILES,weighted_quantiles(y,w))})
    out.update({f'abs_q{int(q*100):02d}':float(v) for q,v in zip(PERCENTILES,weighted_quantiles(np.abs(y),w))})
    return out

def interpolated(t,y,target,max_gap,period=None,discrete=False):
    good=np.isfinite(y)&np.isfinite(t)
    tx=np.asarray(t)[good];yy=np.asarray(y)[good]
    out=np.full(len(target),np.nan)
    if len(tx)<2:return out
    if period:yy=np.unwrap(yy*2*np.pi/period)*period/(2*np.pi)
    j=np.searchsorted(tx,target,side='right')-1;j=np.clip(j,0,len(tx)-2)
    allow=(target>=tx[0])&(target<=tx[-1])&((tx[j+1]-tx[j])<=max_gap)
    out[allow]=yy[j[allow]] if discrete else np.interp(target[allow],tx,yy)
    # Preserve actual support points even if an adjacent gap is long.
    index=np.searchsorted(tx,target);valid=index<len(tx)
    exact=np.zeros(len(target),bool);exact[valid]=np.isclose(tx[index[valid]],target[valid],atol=1e-10,rtol=0)
    out[exact]=yy[index[exact]]
    return out

def finite_blocks(mask):
    padded=np.r_[False,mask,False].astype(int);return zip(np.flatnonzero(np.diff(padded)==1),np.flatnonzero(np.diff(padded)==-1))

def smooth(y,n):
    out=np.full(len(y),np.nan)
    for lo,hi in finite_blocks(np.isfinite(y)):
        # A shorter asymmetric window at an edge creates a false position shift
        # and false derivative spikes. Leave unsupported edge samples missing.
        out[lo:hi]=pd.Series(y[lo:hi]).rolling(n,center=True,min_periods=n).mean().to_numpy()
    return out

def derivative(y,t):
    out=np.full(len(y),np.nan)
    for lo,hi in finite_blocks(np.isfinite(y)):
        if hi-lo>=3:out[lo:hi]=np.gradient(y[lo:hi],t[lo:hi],edge_order=2)
    return out

def statistic_unit(stat,source_unit):
    if stat in {'finite_fraction','coverage_fraction','circular_resultant','active_fraction'}:return 'ratio'
    if stat in {'unique_values','episode_count','reversal_count'}:return 'count'
    if stat in {'valid_duration_s','active_duration_s','longest_episode_s','median_episode_s'}:return 's'
    if stat in {'time_integral','abs_time_integral'}:return '('+source_unit+')*s'
    if stat=='variation_per_s':return '('+source_unit+')/s'
    return source_unit

def spectrum(y,hz):
    """Descriptive spectrum of the longest valid block; no spectral inference."""
    blocks=list(finite_blocks(np.isfinite(y)))
    if not blocks:return {}
    lo,hi=max(blocks,key=lambda v:v[1]-v[0]);z=y[lo:hi].copy()
    if len(z)<max(16,int(2*hz)):return {}
    z-=np.mean(z);freq=np.fft.rfftfreq(len(z),1/hz)
    power=np.abs(np.fft.rfft(z*np.hanning(len(z))))**2;power[0]=0
    total=float(power.sum())
    out={'block_duration_s':len(z)/hz,'frequency_resolution_hz':hz/len(z)}
    if total<1e-20:return out
    out['dominant_frequency_hz']=float(freq[np.argmax(power)])
    out['spectral_centroid_hz']=float(np.dot(freq,power)/total)
    for left,right in [(0,.5),(.5,2),(2,5),(5,10),(10,hz/2)]:
        if right>left:out[f'power_fraction_{left:g}_{right:g}_hz']=float(power[(freq>left)&(freq<=right)].sum()/total)
    return out

def unit(c,a):
    if c in ['x','y','z']:return 'm'
    if c=='speedKmh':return 'km/h'
    if c in ['gas','brake']:return 'fraction (recorded convention)'
    if c in ATTITUDE:return a.attitude_unit
    if c=='steerAngle':return a.steer_unit
    if c.startswith('SlipAngle_'):return a.slip_angle_unit
    return 'native/unverified'

def derived_series(lap,a):
    t=lap.t_sec.to_numpy(float);n=int(np.floor((t[-1]-t[0])*a.grid_hz))+1
    grid=t[0]+np.arange(n)/a.grid_hz
    signals={};units={};notes={};win=max(1,int(round(a.smooth_seconds*a.grid_hz)))
    if win%2==0:win+=1
    def add(c,y,u,note):signals[c]=np.asarray(y,float);units[c]=u;notes[c]=note
    for c in lap:
        if c in META or c.startswith('_') or not np.isfinite(lap[c]).any():continue
        period=(360 if a.attitude_unit=='deg' else 2*np.pi) if c in ATTITUDE else None
        v=interpolated(t,lap[c].to_numpy(float),grid,a.max_gap_seconds,period,c in DISCRETE)
        add(c,smooth(v,win) if c not in DISCRETE else v,unit(c,a),'Uniform-time resampling; centered smoothing; no gap bridging')
    if 'speedKmh' in signals:
        v=signals['speedKmh']/3.6
        add('speed_m_s',v,'m/s','Recorded speed / 3.6')
        acc=derivative(v,grid)
        add('speed_derivative_m_s2',acc,'m/s^2','Speed derivative; not body-frame longitudinal acceleration')
        add('speed_second_derivative_m_s3',derivative(acc,grid),'m/s^3','Derivative of the speed derivative; smoothing-dependent')
    for c in ['gas','brake','steerAngle']+list(ATTITUDE):
        if c in signals:
            add(c+'_rate',derivative(signals[c],grid),units[c]+'/s','Rate after declared uniform-time smoothing')
    if 'heading_rate' in signals and 'speed_m_s' in signals:
        yaw=signals['heading_rate']*(np.pi/180 if a.attitude_unit=='deg' else 1)
        add('yaw_rate_rad_s',yaw,'rad/s','Heading derivative; not an independently measured yaw-rate sensor')
        add('v_times_yaw_rate_m_s2',signals['speed_m_s']*yaw,'m/s^2','Kinematic lateral-acceleration proxy; sideslip/axis assumptions not validated')
    if {'x','z'}.issubset(signals):
        vx=derivative(signals['x'],grid);vz=derivative(signals['z'],grid)
        ax=derivative(vx,grid);az=derivative(vz,grid);v=np.hypot(vx,vz)
        safe=np.where(v>.5,v,np.nan) # Numerical low-speed mask, not a safety limit.
        add('xz_speed_m_s',v,'m/s','Horizontal position derivative')
        add('xz_acceleration_norm_m_s2',np.hypot(ax,az),'m/s^2','Horizontal trajectory acceleration magnitude')
        add('path_curvature_signed_m_inv',(vx*az-vz*ax)/safe**3,'1/m','Driven X-Z path curvature, NOT road-centerline curvature')
        add('path_curvature_abs_m_inv',np.abs(signals['path_curvature_signed_m_inv']),'1/m','Absolute driven-path curvature')
        add('path_lateral_accel_m_s2',(vx*az-vz*ax)/safe,'m/s^2','Signed normal acceleration of horizontal driven trajectory')
        add('path_tangential_accel_m_s2',(vx*ax+vz*az)/safe,'m/s^2','Tangential acceleration of horizontal driven trajectory')
        if 'speed_m_s' in signals:add('position_vs_recorded_speed_m_s',v-signals['speed_m_s'],'m/s','Horizontal derivative speed minus recorded speed; diagnostic')
        if 'y' in signals:
            vy=derivative(signals['y'],grid)
            add('vertical_velocity_m_s',vy,'m/s','Vertical trajectory velocity')
            add('vertical_acceleration_m_s2',derivative(vy,grid),'m/s^2','Vertical trajectory acceleration')
            add('trajectory_grade_proxy',vy/safe,'ratio','Driven trajectory dy/ds_horizontal, not surveyed road grade')
    for family in ['SlipAngle','wheelSlip','NdSlip']:
        cols=[family+'_'+w for w in WHEELS]
        if not all(c in signals for c in cols):continue
        ar=np.column_stack([signals[c] for c in cols]);u=units[cols[0]]
        for suffix,val in {'four_wheel_mean':ar.mean(1),'four_wheel_abs_mean':np.abs(ar).mean(1),
            'four_wheel_abs_max':np.abs(ar).max(1),'four_wheel_spread':np.ptp(ar,axis=1),
            'front_mean':ar[:,:2].mean(1),'rear_mean':ar[:,2:].mean(1),
            'rear_minus_front':ar[:,2:].mean(1)-ar[:,:2].mean(1),
            'front_left_minus_right':ar[:,0]-ar[:,1],'rear_left_minus_right':ar[:,2]-ar[:,3]}.items():
            add(family+'_'+suffix,val,u,'Algebraic wheel-channel comparison; not a tire-force/understeer diagnosis')
    for w in WHEELS:
        if {'NdSlip_'+w,'wheelSlip_'+w}.issubset(signals):
            add('NdSlip_minus_wheelSlip_'+w,signals['NdSlip_'+w]-signals['wheelSlip_'+w],'native/unverified','Difference between recorded channels; not assumed equivalent or perfectly synchronized')
    if {'gas','brake'}.issubset(signals):add('gas_times_brake',signals['gas']*signals['brake'],'fraction^2','Continuous simultaneous-input descriptor')
    return grid,signals,units,notes,win

def event_rows(grid,values,threshold,mode,channel):
    # Midpoint interval classification; first/last runs are explicitly censored.
    mid=(values[:-1]+values[1:])/2;valid=np.isfinite(values[:-1])&np.isfinite(values[1:])
    mask=(np.abs(mid)>threshold if mode=='abs_gt' else mid>threshold if mode=='gt' else mid<threshold)&valid
    dt=np.diff(grid);events=[]
    for lo,hi in finite_blocks(mask):
        events.append({'channel':channel,'operator':mode,'threshold':threshold,'start_t_sec':float(grid[lo]),'end_t_sec':float(grid[hi]),
            'duration_s':float(grid[hi]-grid[lo]),'left_censored':lo==0 or not valid[lo-1],
            'right_censored':hi==len(mask) or (hi<len(mask) and not valid[hi])})
    valid_s=float(dt[valid].sum());active_s=float(dt[mask].sum())
    summary={'valid_duration_s':valid_s,'active_duration_s':active_s,'active_fraction':active_s/valid_s if valid_s else np.nan,
             'episode_count':float(len(events)),'longest_episode_s':max([e['duration_s'] for e in events],default=0.),
             'median_episode_s':float(np.median([e['duration_s'] for e in events])) if events else 0.}
    return summary,events

def spatial_profile(lap,grid,signals,a):
    p=lap['_progress'].to_numpy(float);t=lap.t_sec.to_numpy(float)
    # Use first-forward-passage samples; do not turn reverse motion into extra laps.
    keep=np.r_[True,np.diff(np.maximum.accumulate(p))>1e-12];pp=p[keep];tt=t[keep]
    phase=np.arange(a.profile_points)/a.profile_points
    direction=1 if a.direction=='forward' else -1
    target=np.mod(direction*(phase-float(lap.normalizedCarPosition.iloc[0])),1.)
    target_t=interpolated(pp,tt,target,a.max_progress_gap)
    out={'track_progress':phase}
    for c in signals:
        if c in lap:
            period=(360 if a.attitude_unit=='deg' else 2*np.pi) if c in ATTITUDE else None
            out[c]=interpolated(t,lap[c].to_numpy(float),target_t,a.max_gap_seconds,period,c in DISCRETE)
        else:
            out[c]=interpolated(grid,signals[c],target_t,a.max_gap_seconds,discrete=c in DISCRETE)
        if c in ATTITUDE:
            period=360 if a.attitude_unit=='deg' else 2*np.pi
            out[c]=(out[c]+period/2)%period-period/2
    return pd.DataFrame(out)

def process(info,a):
    lap,qc,inventory=read_and_extract(info['source'],a)
    t=lap.t_sec.to_numpy(float);features=[];catalog=[];events=[];correlations=[]
    def feature(channel,stat,value,u,stage):
        features.append({**info,'stage':stage,'channel':channel,'statistic':stat,'value':value,'unit':u})
    feature('lap','duration_s',qc['lap_time_s'],'s','lap')
    feature('lap','legacy_sampled_duration_s',qc['sampled_crossing_lap_time_s'],'s','lap')
    for c in lap:
        if c in META or c.startswith('_'):continue
        period=(360 if a.attitude_unit=='deg' else 2*np.pi) if c in ATTITUDE else None
        catalog.append({'channel':c,'unit':unit(c,a),'stage':'raw','definition':'Recorded channel; physical meaning beyond name not inferred'})
        for stat,val in summarize(t,lap[c].to_numpy(float),a.max_gap_seconds,period).items():feature(c,stat,val,statistic_unit(stat,unit(c,a)),'raw')
    # Length estimates retain position jumps as flagged diagnostics; no silent removal.
    for name,cols in [('path_length_3d_m',['x','y','z']),('path_length_xz_m',['x','z'])]:
        if all(c in lap for c in cols):
            d=np.linalg.norm(np.diff(lap[cols].to_numpy(float),axis=0),axis=1)
            good=np.isfinite(d)&(np.diff(t)<=a.max_gap_seconds)
            feature('lap',name,float(d[good].sum()),'m','lap')
    if 'speedKmh' in lap:
        q=summarize(t,lap.speedKmh.to_numpy(float)/3.6,a.max_gap_seconds)
        if 'time_integral' in q:feature('lap','speed_integrated_distance_m',q['time_integral'],'m','lap')
    grid,signals,units,notes,win=derived_series(lap,a);qc['smoothing_samples']=win
    for c,y in signals.items():
        period=(360 if a.attitude_unit=='deg' else 2*np.pi) if c in ATTITUDE else None
        catalog.append({'channel':c,'unit':units[c],'stage':'regularized','definition':notes[c]})
        for stat,val in summarize(grid,y,a.max_gap_seconds,period).items():feature(c,stat,val,statistic_unit(stat,units[c]),'regularized')
    for c in ['gas','brake','steerAngle','speedKmh','yaw_rate_rad_s','SlipAngle_four_wheel_abs_mean']:
        if c not in signals:continue
        for stat,val in spectrum(signals[c],a.grid_hz).items():
            feature(c,stat,val,'ratio' if stat.startswith('power_fraction') else 's' if stat.endswith('_s') else 'Hz','spectrum')
    if 'steerAngle' in signals:
        y=signals['steerAngle']
        for deadband in [0.,.001,.01]:
            count=0
            for lo,hi in finite_blocks(np.isfinite(y)):
                signs=np.sign(y[lo:hi][np.abs(y[lo:hi])>deadband])
                count+=int(np.sum(np.diff(signs)!=0))
            feature('steerAngle',f'reversals_deadband_{deadband:g}',count,'count','diagnostic_reversals')
    thresholds={'gas':[(v,'gt') for v in [.05,.1,.5,.8,.95]],'brake':[(v,'gt') for v in [.01,.05,.1,.2,.5]],
        'speed_derivative_m_s2':[(v,'gt') for v in [0,1,2,3]]+[(v,'lt') for v in [-1,-2,-3]],
        'speedKmh':[(v,'lt') for v in [1,5,30,60,100]]}
    if a.slip_angle_unit!='native':
        factor=1 if a.slip_angle_unit=='deg' else np.pi/180
        for c in signals:
            if c.startswith('SlipAngle_'):thresholds[c]=[(v*factor,'abs_gt') for v in [1,2,3,5,10,15]]
    for c,specs in thresholds.items():
        if c not in signals:continue
        for threshold,mode in specs:
            stats,episodes=event_rows(grid,signals[c],threshold,mode,c)
            for stat,val in stats.items():feature(c,f'{mode}_{threshold:g}__{stat}',val,statistic_unit(stat,units[c]),'events')
            events.extend([{**info,**e,'threshold_unit':units[c]} for e in episodes])
    if {'gas','brake'}.issubset(signals):
        for threshold in [.01,.05,.1]:
            stats,episodes=event_rows(grid,np.minimum(signals['gas'],signals['brake']),threshold,'gt','gas_and_brake')
            for stat,val in stats.items():feature('gas_and_brake',f'gt_{threshold:g}__{stat}',val,statistic_unit(stat,'fraction'),'events')
            events.extend([{**info,**e,'threshold_unit':'fraction'} for e in episodes])
    # Pairwise association only, no lag/causal/control-mechanism interpretation.
    candidates=[c for c in ['speedKmh','gas','brake','steerAngle','speed_derivative_m_s2','yaw_rate_rad_s','path_curvature_abs_m_inv','SlipAngle_four_wheel_abs_mean','SlipAngle_rear_minus_front','NdSlip_four_wheel_abs_mean','wheelSlip_four_wheel_abs_mean'] if c in signals]
    for x,y in itertools.combinations(candidates,2):
        mask=np.isfinite(signals[x])&np.isfinite(signals[y]);xx=signals[x][mask];yy=signals[y][mask]
        if len(xx)>3 and np.std(xx)>1e-12 and np.std(yy)>1e-12:
            correlations.append({**info,'x':x,'y':y,'n_grid_samples':len(xx),'pearson_r':float(np.corrcoef(xx,yy)[0,1]),
                'spearman_r':float(pd.Series(xx).rank().corr(pd.Series(yy).rank()))})
    profile=spatial_profile(lap,grid,signals,a)
    for k,v in info.items():profile.insert(0,k,v)
    return features,{**info,**qc},[{**info,**r} for r in inventory],catalog,events,correlations,profile,lap,pd.DataFrame({'t_sec':grid,**signals})

def aggregate(features,profiles,inputs,output,a):
    f=pd.DataFrame(features);keys=['config','track','vehicle','setting','stage','channel','statistic','unit']
    summary=f.groupby(keys,dropna=False)['value'].agg(n='count',mean='mean',sample_sd='std',median='median',min='min',max='max').reset_index()
    summary['interpretation']='Across-run descriptive statistics; n is finite contributing runs, not telemetry samples'
    write_csv(summary,output/'condition_summary.csv')
    base=summary[summary.setting==80].drop(columns='setting')
    join=['config','track','vehicle','stage','channel','statistic','unit']
    contrast=summary.merge(base[join+['mean','n']],on=join,how='left',suffixes=('','_at_80'))
    contrast['difference_from_80']=contrast['mean']-contrast['mean_at_80']
    # Ratios only for lap time; avoid invalid ratios of signed/angle features.
    ok=(contrast.channel=='lap')&(contrast.statistic=='duration_s')&(contrast.mean_at_80>0)
    contrast['lap_time_percent_change_from_80']=np.where(ok,100*contrast.difference_from_80/contrast.mean_at_80,np.nan)
    write_csv(contrast,output/'aipush_contrasts.csv')
    comparisons=[]
    for axis,fixed in [('vehicle','track'),('track','vehicle')]:
        groupkeys=[fixed,'setting','stage','channel','statistic','unit']
        for key,g in summary.groupby(groupkeys,dropna=False):
            for (_,left),(_,right) in itertools.combinations(g.sort_values(axis).iterrows(),2):
                comparisons.append({**dict(zip(groupkeys,key)),'contrast_axis':axis,'level_a':left[axis],'level_b':right[axis],
                    'mean_a':left['mean'],'mean_b':right['mean'],'difference_b_minus_a':right['mean']-left['mean'],'n_a':left['n'],'n_b':right['n']})
    write_csv(comparisons,output/'vehicle_track_contrasts.csv')
    # Profile means/SD across runs at an absolute track location; not a model envelope.
    p=pd.concat(profiles,ignore_index=True)
    idcols=set(inputs[0])|{'track_progress'}
    numeric=[c for c in p if c not in idcols]
    out=[]
    for key,g in p.groupby(['config','track','vehicle','setting'],sort=True):
        meta=dict(zip(['config','track','vehicle','setting'],key))
        for c in numeric:
            gg=g[['track_progress',c]].copy()
            if c in ATTITUDE:
                period=360 if a.attitude_unit=='deg' else 2*np.pi
                gg['sin']=np.sin(gg[c]*2*np.pi/period);gg['cos']=np.cos(gg[c]*2*np.pi/period)
                s=gg.groupby('track_progress').agg(n=(c,'count'),sin=('sin','mean'),cos=('cos','mean')).reset_index()
                s['mean']=np.arctan2(s['sin'],s['cos'])*period/(2*np.pi)
                s['sample_sd']=np.nan
                s['circular_sd']=np.sqrt(-2*np.log(np.hypot(s['sin'],s['cos']).clip(1e-15,1)))*period/(2*np.pi)
                s=s.drop(columns=['sin','cos'])
            else:
                s=gg.groupby('track_progress')[c].agg(n='count',mean='mean',sample_sd='std',min='min',max='max').reset_index()
            for k,v in meta.items():s[k]=v
            s['channel']=c;out.append(s)
    write_csv(pd.concat(out,ignore_index=True),output/'condition_profiles.csv.gz')

def self_test():
    import importlib.util
    path=Path(__file__).with_name('test_ai_preprocessing.py')
    spec=importlib.util.spec_from_file_location('test_ai_preprocessing',path);mod=importlib.util.module_from_spec(spec);spec.loader.exec_module(mod)
    import unittest
    result=unittest.TextTestRunner(verbosity=2).run(unittest.defaultTestLoader.loadTestsFromModule(mod))
    if not result.wasSuccessful():raise SystemExit(1)

def main(argv=None):
    a=arguments(argv)
    if a.self_test:self_test();return
    if a.grid_hz<2 or a.profile_points<10 or a.max_gap_seconds<=0 or a.smooth_seconds<0 or a.max_progress_gap<=0 or not 0<a.max_progress_step<.5:raise ValueError('Invalid grid/gap/smoothing/progress options')
    if a.limit is not None and a.limit<1:raise ValueError('--limit must be positive')
    inputs=discover(a)
    if not inputs:raise FileNotFoundError('No *+AI/<setting>/*.xlsx found')
    if a.inspect:inspect(a,inputs);return
    selected=inputs[:a.limit] if a.limit else inputs
    output=(a.output or a.root/'ai_driving_characteristics').resolve()
    for r in selected:
        source=Path(r['source']).resolve()
        if output==source.parent or output in source.parents:raise ValueError('Output must not be an input ancestor/directory')
    output.mkdir(parents=True,exist_ok=True)
    feature_rows=[];qc_rows=[];inventory_rows=[];catalog=[];correlation_rows=[];profiles=[];failures=[];provenance=[];success=[]
    all_keys=['config','track','vehicle','setting','run','source']
    write_json({'status':'running','version':VERSION,'arguments':vars(a),'discovered':len(inputs),'selected':len(selected),'results_authority':'Only run_manifest.json with status complete or partial after this invocation is authoritative'},output/'run_manifest.json')
    for index,info in enumerate(selected,1):
        print(f"[{index}/{len(selected)}] {info['config']} AIpush={info['setting']} run={info['run']}",flush=True)
        path=Path(info['source']);entry={**info,'size':path.stat().st_size,'mtime_ns':path.stat().st_mtime_ns}
        try:
            if not a.no_hash:entry['sha256']=sha(path)
            features,qc,inventory,channel_defs,events,corr,profile,lap,ts=process(info,a)
            stem=output/'runs'/info['config']/str(info['setting'])/info['run']
            write_csv(features,stem/'features.csv');write_csv(events,stem/'diagnostic_episodes.csv.gz')
            write_csv(profile,stem/'profile.csv.gz')
            if a.save_clean_laps:write_csv(lap,stem/'complete_lap.csv.gz')
            if a.save_time_series:write_csv(ts,stem/'regularized_timeseries.csv.gz')
            feature_rows.extend(features);qc_rows.append(qc);inventory_rows.extend(inventory);catalog.extend(channel_defs);correlation_rows.extend(corr);profiles.append(profile);success.append(info)
            entry['status']='ok'
        except Exception as exc:
            entry['status']='failed';entry['error']=str(exc)
            failures.append({**info,'error':str(exc),'traceback':traceback.format_exc()})
            print('  FAILED: '+str(exc),flush=True)
            if a.fail_fast:
                write_json({'status':'failed','failure':entry},output/'run_manifest.json');raise
        provenance.append(entry)
        write_csv(qc_rows,output/'lap_qc.csv');write_csv(pd.DataFrame(failures,columns=all_keys+['error','traceback']),output/'failures.csv')
    if feature_rows:
        write_csv(feature_rows,output/'run_features_long.csv');write_csv(inventory_rows,output/'channel_inventory.csv')
        write_csv(pd.DataFrame(catalog).drop_duplicates(),output/'channel_dictionary.csv')
        write_csv(correlation_rows,output/'within_run_correlations.csv')
        aggregate(feature_rows,profiles,success,output,a)
    else:
        # Invalidate aggregate products from any earlier successful invocation.
        for filename in ['run_features_long.csv','channel_inventory.csv','channel_dictionary.csv','within_run_correlations.csv','condition_summary.csv','aipush_contrasts.csv','vehicle_track_contrasts.csv','condition_profiles.csv.gz']:
            write_csv(pd.DataFrame(columns=['no_successful_runs_this_invocation']),output/filename)
    coverage=[]
    successful={(r['config'],r['setting'],r['run']) for r in success}
    for config in sorted({r['config'] for r in inputs}):
        for setting in a.settings:
            n=sum(r['config']==config and r['setting']==setting for r in inputs)
            passed=sum(k[0]==config and k[1]==setting for k in successful)
            coverage.append({'config':config,'setting':setting,'discovered_runs':n,'processed_ok':passed,'expected_repeats':a.expected_repeats,'complete_expected_count':passed==a.expected_repeats})
    write_csv(coverage,output/'design_coverage.csv')
    design_complete=len({r['config'] for r in inputs})==a.expected_configs and all(r['complete_expected_count'] for r in coverage)
    write_json({'status':'complete' if not failures and len(selected)==len(inputs) and design_complete else 'partial','expected_design_complete':design_complete,'version':VERSION,'arguments':vars(a),'inputs':provenance,'success':len(success),'failed':len(failures),
        'discovered':len(inputs),'selected':len(selected),'script_sha256':sha(__file__),
        'dependencies':{'numpy':np.__version__,'pandas':pd.__version__},
        'lap_rule':'First recorded phase to exactly +1 directed unwrapped progress; endpoint interpolated. Legacy first-crossing duration also exported.',
        'statistics':'Trapezoidal time-weighted moments and endpoint-weighted quantiles; raw and smoothed stages kept separate. Across-run SD uses ddof=1.',
        'thresholds':'Descriptive diagnostic grids only, NOT crash, drift, accuracy, or safety limits.',
        'profiles':'Absolute normalized track position; first-forward-passage samples. Recorded channels interpolated directly from raw lap; derived channels from the regularized stage. No interpolation across declared progress/time gaps.',
        'output_authority':'Aggregate files and listed successful runs from this manifest only. Older per-run files may remain after a reduced/failed rerun; do not glob them as current evidence.',
        'limitations':['Observational trajectory descriptors, not intrinsic vehicle or surveyed track properties.','No engine/suspension/load signals are invented; inventory states what was recorded.','Native wheelSlip/NdSlip/steer units unverified unless explicitly configured.','Start state varies; repeated runs are not paired trajectories.','Derivatives and episodes depend on resampling/smoothing; rerun sensitivity with separate --output directories.']},output/'run_manifest.json')
    print(f"Done: {len(success)} successful, {len(failures)} failed. Results: {output}")
    if failures or not success:raise SystemExit(2)

if __name__=='__main__':main()
